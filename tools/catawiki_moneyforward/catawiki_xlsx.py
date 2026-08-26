# -*- coding: utf-8 -*-
import pickle, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

detail, excluded, agg, journals = pickle.load(open('catawiki.pkl','rb'))
fee_agg = pickle.load(open('catawiki_fee.pkl','rb'))
import csv as _csv
from decimal import Decimal, ROUND_HALF_UP
_rates = {r['date']: Decimal(r['eurjpy']) for r in _csv.DictReader(open('catawiki_rates.csv'))}
def _yen(x): return int(x.quantize(Decimal('1'), rounding=ROUND_HALF_UP))
RATE_URL = 'https://docs.google.com/spreadsheets/d/1NEvV57OcsDs5qVpBEB2g-FezXJydnUpvt0CdMT_b6Z4/edit'
wb = Workbook()
HDR = PatternFill('solid', fgColor='1F3864'); HDRF = Font(color='FFFFFF', bold=True); TITLE = Font(bold=True, size=13)

def head(ws, cols, row=1):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.fill, cell.font = HDR, HDRF
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
def widths(ws, w):
    for i, x in enumerate(w, 1): ws.column_dimensions[get_column_letter(i)].width = x

ws = wb.active; ws.title = '概要'
lines = [('catawiki売上 マネーフォワード取込資料', TITLE),
 ('ワールドドリーム株式会社／会計期間 2025年7月1日〜2026年6月30日', None), ('', None),
 ('■ 集計対象', TITLE),
 ('catawiki Orders レポート（20件）のうち、キャンセル済み2件を除く18件', None),
 ('計上日：Order date（落札・注文成立日）', None), ('', None),
 ('■ 計上金額', TITLE),
 ('Highest bid（落札額）＋ Shipping costs（送料）の総額', None),
 ('Commission（catawiki手数料）は控除せず、eBayと同様に別資料で費用計上', None), ('', None),
 ('■ 円換算', TITLE),
 ('取引ごとにOrder dateのEUR/JPYレートで換算（円未満四捨五入）', None),
 ('レート出典：Google Finance（GOOGLEFINANCE関数・日次終値）', None),
 (f'レート原本：{RATE_URL}', None), ('', None),
 ('■ 仕訳', TITLE),
 ('月次で1本に集約し、各月末日付で起票', None),
 ('売上：借）売掛金［catawiki］対象外 ／ 貸）売上高［catawiki］輸出売上 0%', None),
 ('手数料：借）支払手数料［catawiki］対象外 ／ 貸）売掛金［catawiki］対象外', None),
 ('※支払手数料の補助科目catawikiに設定された既定の税区分（対象外）に従っています。', None),
 ('　eBay（課税仕入 10%）とは扱いが異なります。', None), ('', None),
 ('■ 集計結果', TITLE)]
r = 1
for t, f in lines:
    c = ws.cell(row=r, column=1, value=t)
    if f: c.font = f
    r += 1
head(ws, ['月','件数','売上(EUR)','円換算額','（参考）手数料(EUR)'], row=r); r += 1
for ym, a in sorted(agg.items()):
    ws.cell(row=r, column=1, value=ym); ws.cell(row=r, column=2, value=a['n'])
    ws.cell(row=r, column=3, value=float(a['fx'])).number_format = '#,##0.00'
    ws.cell(row=r, column=4, value=a['jpy']).number_format = '#,##0'
    ws.cell(row=r, column=5, value=float(a['fee'])).number_format = '#,##0.00'
    r += 1
ws.cell(row=r, column=1, value='合計').font = Font(bold=True)
ws.cell(row=r, column=2, value=sum(a['n'] for a in agg.values())).font = Font(bold=True)
c = ws.cell(row=r, column=3, value=float(sum(a['fx'] for a in agg.values()))); c.number_format='#,##0.00'; c.font=Font(bold=True)
c = ws.cell(row=r, column=4, value=sum(a['jpy'] for a in agg.values())); c.number_format='#,##0'; c.font=Font(bold=True)
c = ws.cell(row=r, column=5, value=float(sum(a['fee'] for a in agg.values()))); c.number_format='#,##0.00'; c.font=Font(bold=True)
widths(ws, [26,10,14,14,20])

ws = wb.create_sheet('月次仕訳(MF取込内容)')
head(ws, ['取引No','取引日','借方勘定科目','借方補助科目','借方税区分','借方金額(円)','貸方勘定科目','貸方補助科目','貸方税区分','貸方金額(円)','摘要'])
for i, j in enumerate(journals, 1):
    ws.append([i, j['取引日'], j['借方勘定科目'], j['借方補助科目'], j['借方税区分'], j['借方金額'],
               j['貸方勘定科目'], j['貸方補助科目'], j['貸方税区分'], j['貸方金額'], j['摘要']])
import calendar as _cal, datetime as _dt
for k, (ym, a) in enumerate(sorted(fee_agg.items()), len(journals) + 1):
    y, m = int(ym[:4]), int(ym[5:])
    ws.append([k, _dt.date(y, m, _cal.monthrange(y, m)[1]).strftime('%Y/%m/%d'), '支払手数料', 'catawiki', '対象外', a['jpy'],
               '売掛金', 'catawiki', '対象外', a['jpy'], f'catawiki手数料 {y}年{m}月（{a["n"]}件 EUR {a["fx"]}）'])
for col in (6, 10):
    for row in ws.iter_rows(min_row=2, min_col=col, max_col=col): row[0].number_format = '#,##0'
widths(ws, [8,12,14,16,18,14,14,16,18,14,54])

ws = wb.create_sheet('取引明細')
cols = ['日付','注文番号','品名','国','通貨','落札額','送料','売上額(外貨)','適用レート','レート基準日','円換算額','（参考）手数料(外貨)','入金予定日','請求書番号','レート出典']
head(ws, cols)
for d in sorted(detail, key=lambda x: x['日付']):
    ws.append([d['日付'], d['注文番号'], d['品名'], d['国'], d['通貨'], d['落札額'], d['送料'], d['売上額(外貨)'],
               d['適用レート'], d['レート基準日'], d['円換算額'], d['手数料(外貨)'], d['入金予定日'], d['請求書番号'], d['レート出典']])
for col in (6, 7, 8, 12):
    for row in ws.iter_rows(min_row=2, min_col=col, max_col=col): row[0].number_format = '#,##0.00'
for row in ws.iter_rows(min_row=2, min_col=9, max_col=9): row[0].number_format = '#,##0.0000'
for row in ws.iter_rows(min_row=2, min_col=11, max_col=11): row[0].number_format = '#,##0'
widths(ws, [12,12,40,16,8,10,10,14,12,12,12,18,12,16,60])
ws.auto_filter.ref = ws.dimensions

ws = wb.create_sheet('手数料明細')
ws.cell(row=1, column=1, value='catawiki手数料（Commission incl. VAT）の円換算明細').font = TITLE
ws.cell(row=2, column=1, value='売上と同じくOrder dateのEUR/JPYレートで1件ずつ換算し、月次で支払手数料に計上')
head(ws, ['日付','注文番号','手数料(EUR)','適用レート','円換算額'], row=4)
for d in sorted(detail, key=lambda x: x['日付']):
    if d['手数料(外貨)'] == 0: continue
    f = Decimal(str(d['手数料(外貨)'])); rt = _rates[d['レート基準日']]
    ws.append([d['日付'], d['注文番号'], float(f), float(rt), _yen(f * rt)])
r2 = ws.max_row + 1
ws.cell(row=r2, column=1, value='合計').font = Font(bold=True)
c = ws.cell(row=r2, column=3, value=float(sum(a['fx'] for a in fee_agg.values()))); c.number_format='#,##0.00'; c.font=Font(bold=True)
c = ws.cell(row=r2, column=5, value=sum(a['jpy'] for a in fee_agg.values())); c.number_format='#,##0'; c.font=Font(bold=True)
for row in ws.iter_rows(min_row=5, min_col=3, max_col=3): row[0].number_format = '#,##0.00'
for row in ws.iter_rows(min_row=5, min_col=4, max_col=4): row[0].number_format = '#,##0.0000'
for row in ws.iter_rows(min_row=5, min_col=5, max_col=5): row[0].number_format = '#,##0'
widths(ws, [12,14,14,12,12])

ws = wb.create_sheet('除外取引')
ws.cell(row=1, column=1, value='キャンセル等により売上計上から除外した取引').font = TITLE
head(ws, ['日付','注文番号','品名','落札額','送料','手数料','請求書番号','除外理由'], row=3)
for e in excluded:
    ws.append([e['日付'], e['注文番号'], e['品名'], e['落札額'], e['送料'], e['手数料'], e['請求書番号'], e['除外理由']])
widths(ws, [12,12,40,10,10,10,16,60])

ws = wb.create_sheet('適用為替レート')
ws.cell(row=1, column=1, value='本資料で適用したEUR/JPYレート').font = TITLE
ws.cell(row=2, column=1, value='出典：Google Finance（GOOGLEFINANCE関数・日次終値）')
ws.cell(row=3, column=1, value=f'レート原本：{RATE_URL}')
head(ws, ['レート基準日','通貨','対円レート','適用件数'], row=5)
cnt = collections.Counter(d['レート基準日'] for d in detail)
seen = {}
for d in detail: seen[d['レート基準日']] = d['適用レート']
for dt in sorted(seen):
    ws.append([dt, 'EUR', seen[dt], cnt[dt]])
for row in ws.iter_rows(min_row=6, min_col=3, max_col=3): row[0].number_format = '#,##0.0000'
widths(ws, [14,8,14,10])

out = 'catawiki売上_月次集計_円換算_2026年5月-6月.xlsx'
wb.save(out); print('saved', out)
