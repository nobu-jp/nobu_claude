# -*- coding: utf-8 -*-
import pickle, collections, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

detail, excluded, agg, journals, applied = pickle.load(open('data.pkl','rb'))
RATE_SHEET_URL = 'https://docs.google.com/spreadsheets/d/14ms_NDm2ty8A0jLRwgnAij2aYVwAx4O0bHAFDiGNNZg/edit'

wb = Workbook()
HDR = PatternFill('solid', fgColor='1F3864')
HDRF = Font(color='FFFFFF', bold=True)
TITLE = Font(bold=True, size=13)
THIN = Border(*[Side(style='thin', color='BFBFBF')]*4)

def head(ws, cols, row=1):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.fill, cell.font = HDR, HDRF
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.freeze_panes = ws.cell(row=row+1, column=1)

def widths(ws, w):
    for i, x in enumerate(w, 1): ws.column_dimensions[get_column_letter(i)].width = x

# ---------- 1. 概要 ----------
ws = wb.active; ws.title = '概要'
lines = [
 ('eBay売上 マネーフォワード取込資料', TITLE),
 ('ワールドドリーム株式会社／会計期間 2025年7月1日〜2026年6月30日', None), ('', None),
 ('■ 集計対象', TITLE),
 ('eBay Transaction report 3アカウント分（go-go-japan／yokoso2020japan2／visit3japan2）', None),
 ('取引種別：Order／Adjustment（税額調整分のみ）／Refund／Claim／Payment dispute', None),
 ('除外：Other fee・Payout・Charge・Hold・Shipping label、手数料関連のAdjustment（別資料で計上）', None),
 ('', None),
 ('■ 計上金額', TITLE),
 ('Gross transaction amount（商品代＋送料の総額／eBay代理徴収税は含まない）を使用', None),
 ('複数商品注文は合計行のみ計上し、明細行（金額欄が空欄の行）は二重計上防止のため除外', None),
 ('', None),
 ('■ 円換算', TITLE),
 ('取引ごとに Transaction creation date のレートで換算（端数は円未満四捨五入）', None),
 (f'レート出典：Google Finance（GOOGLEFINANCE関数・日次終値）', None),
 (f'レート原本：{RATE_SHEET_URL}', None),
 ('※銀行公示のTTM（仲値）ではなく市場終値です。継続適用を前提としています。', None),
 ('※AUD建て・EUR建て取引は当該通貨の対円レートで直接換算しています。', None),
 ('', None),
 ('■ 仕訳', TITLE),
 ('月×アカウント×取引種別で1本に集約し、各月末日付で起票', None),
 ('売上計上：借）売掛金［対象外］／貸）売上高［輸出売上 0%］', None),
 ('返金・クレーム・支払異議：借）売上高［輸出売上-返還等 0%］／貸）売掛金［対象外］', None),
 ('補助科目：ebay_go-go-japan／ebay_yokoso2020japan2／ebay_visit3japan2', None),
 ('', None),
 ('■ 集計結果', TITLE),
]
r = 1
for text, f in lines:
    c = ws.cell(row=r, column=1, value=text)
    if f: c.font = f
    r += 1
by_acct = collections.defaultdict(collections.Counter)
for (ym, acct, t), a in agg.items(): by_acct[acct][t] += a['jpy']
types = ['Order','Adjustment','Refund','Claim','Payment dispute']
head(ws, ['アカウント'] + types + ['純額（円）'], row=r); ws.freeze_panes = None
r += 1
for acct in ['go-go-japan','yokoso2020japan2','visit3japan2']:
    ws.cell(row=r, column=1, value=acct)
    for i, t in enumerate(types, 2):
        ws.cell(row=r, column=i, value=by_acct[acct].get(t, 0)).number_format = '#,##0'
    ws.cell(row=r, column=7, value=sum(by_acct[acct].values())).number_format = '#,##0'
    r += 1
ws.cell(row=r, column=1, value='合計').font = Font(bold=True)
for i, t in enumerate(types, 2):
    c = ws.cell(row=r, column=i, value=sum(by_acct[a].get(t, 0) for a in by_acct)); c.number_format='#,##0'; c.font=Font(bold=True)
c = ws.cell(row=r, column=7, value=sum(sum(v.values()) for v in by_acct.values())); c.number_format='#,##0'; c.font=Font(bold=True)
widths(ws, [22,16,14,14,14,16,16])

# ---------- 2. 月次仕訳 ----------
ws = wb.create_sheet('月次仕訳(MF取込内容)')
cols = ['取引No','取引日','借方勘定科目','借方補助科目','借方税区分','借方金額(円)','貸方勘定科目','貸方補助科目','貸方税区分','貸方金額(円)','摘要']
head(ws, cols)
for i, j in enumerate(journals, 1):
    ws.append([i, j['取引日'], j['借方勘定科目'], j['借方補助科目'], j['借方税区分'], j['借方金額'],
               j['貸方勘定科目'], j['貸方補助科目'], j['貸方税区分'], j['貸方金額'], j['摘要']])
for row in ws.iter_rows(min_row=2, min_col=6, max_col=6): row[0].number_format = '#,##0'
for row in ws.iter_rows(min_row=2, min_col=10, max_col=10): row[0].number_format = '#,##0'
widths(ws, [8,12,14,22,20,14,14,22,20,14,70])

# ---------- 3. 月次集計クロス ----------
ws = wb.create_sheet('月次集計')
head(ws, ['月','アカウント','種別','件数','通貨','外貨額','円換算額'])
for (ym, acct, t), a in sorted(agg.items()):
    ws.append([ym, acct, t, a['n'], '/'.join(sorted(a['cur'])), float(a['fx']), a['jpy']])
for row in ws.iter_rows(min_row=2, min_col=6, max_col=6): row[0].number_format = '#,##0.00'
for row in ws.iter_rows(min_row=2, min_col=7, max_col=7): row[0].number_format = '#,##0'
widths(ws, [10,22,18,8,10,14,14])

# ---------- 4. 取引明細 ----------
ws = wb.create_sheet('取引明細')
cols = ['日付','アカウント','種別','計上区分','注文番号','通貨','外貨額','適用レート','レート基準日','円換算額','摘要(eBay側)','レート出典']
head(ws, cols)
for d in sorted(detail, key=lambda x: (x['日付'], x['アカウント'])):
    ws.append([d['日付'], d['アカウント'], d['種別'], d['計上区分'], d['注文番号'], d['通貨'], d['外貨額'],
               d['適用レート'], d['レート基準日'], d['円換算額'], d['摘要元'], d['レート出典']])
for row in ws.iter_rows(min_row=2, min_col=7, max_col=7): row[0].number_format = '#,##0.00'
for row in ws.iter_rows(min_row=2, min_col=8, max_col=8): row[0].number_format = '#,##0.0000'
for row in ws.iter_rows(min_row=2, min_col=10, max_col=10): row[0].number_format = '#,##0'
widths(ws, [12,20,16,12,18,8,12,12,12,12,34,60])
ws.auto_filter.ref = ws.dimensions

# ---------- 5. 除外取引 ----------
ws = wb.create_sheet('除外取引')
ws.cell(row=1, column=1, value='別資料で計上する取引・二重計上防止のため除外した行（サマリ）').font = TITLE
head(ws, ['アカウント','種別','件数','外貨合計(USD等)','除外理由'], row=3)
summ = collections.defaultdict(lambda: [0, 0.0, ''])
for e in excluded:
    k = (e['アカウント'], e['種別'])
    s = summ[k]; s[0] += 1
    if isinstance(e['外貨額'], float): s[1] += e['外貨額']
    s[2] = e['除外理由'].split(':')[0]
r = 4
for (acct, t), (n, fx, why) in sorted(summ.items()):
    ws.append([acct, t, n, round(fx, 2), why]); r += 1
r += 1
ws.cell(row=r, column=1, value='個別確認用：Other fee・Payout 以外の除外行').font = TITLE
r += 1
head(ws, ['日付','アカウント','種別','注文番号','通貨','外貨額','摘要(eBay側)','除外理由'], row=r)
for e in sorted(excluded, key=lambda x: (x['日付'], x['アカウント'])):
    if e['種別'] in ('Other fee', 'Payout'): continue
    r += 1
    for i, v in enumerate([e['日付'], e['アカウント'], e['種別'], e['注文番号'], e['通貨'], e['外貨額'], e['摘要元'], e['除外理由']], 1):
        ws.cell(row=r, column=i, value=v)
widths(ws, [12,20,18,18,8,14,40,52])

# ---------- 6. 適用為替レート ----------
ws = wb.create_sheet('適用為替レート')
ws.cell(row=1, column=1, value='本資料で実際に適用した為替レート一覧').font = TITLE
ws.cell(row=2, column=1, value='出典：Google Finance（GOOGLEFINANCE関数・日次終値）')
ws.cell(row=3, column=1, value=f'レート原本スプレッドシート：{RATE_SHEET_URL}')
ws.cell(row=4, column=1, value='※銀行公示TTM（仲値）ではなく市場終値。継続適用を前提とする。')
head(ws, ['レート基準日','通貨','対円レート','適用件数','出典'], row=6)
cnt = collections.Counter((d['レート基準日'], d['通貨']) for d in detail)
r = 7
for (dt, cur), rate in sorted(applied.items()):
    ws.append([dt, cur, float(rate), cnt[(dt, cur)], 'Google Finance 日次終値'])
    r += 1
for row in ws.iter_rows(min_row=7, min_col=3, max_col=3): row[0].number_format = '#,##0.0000'
widths(ws, [14,8,14,10,28])

out = 'eBay売上_月次集計_円換算_2025年7月-2026年6月.xlsx'
wb.save(out)
print('saved', out, '/ 明細', len(detail), '/ 仕訳', len(journals), '/ 適用レート', len(applied))
